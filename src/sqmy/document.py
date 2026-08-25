from __future__ import annotations

from copy import deepcopy
import os
from pathlib import Path
import shutil
import uuid

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


def _set_font(style, name: str, size: float, bold: bool = False, color: str = "000000") -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)


def _set_run_font(run, name: str) -> None:
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)


def _apply_run_fonts(doc, name: str) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            _set_run_font(run, name)


def _temporary_output(output: Path) -> Path:
    return output.with_name(f".{output.name}.{uuid.uuid4().hex}.tmp")


def _save_atomic(doc, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = _temporary_output(output)
    try:
        doc.save(temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def atomic_copy(source: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = _temporary_output(output)
    try:
        shutil.copy2(source, temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def create_template(reference: Path, output: Path, cfg: dict, signature: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = _temporary_output(output)
    shutil.copy2(reference, temporary)
    doc = Document(temporary)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(cfg["page_width_mm"]), Mm(cfg["page_height_mm"])
    section.top_margin, section.bottom_margin = Mm(cfg["margin_top_mm"]), Mm(cfg["margin_bottom_mm"])
    section.left_margin, section.right_margin = Mm(cfg["margin_left_mm"]), Mm(cfg["margin_right_mm"])
    font = cfg["body_font"]
    _set_font(doc.styles["Normal"], font, cfg["body_size_pt"])
    _set_font(doc.styles["Body Text"], font, cfg["body_size_pt"])
    _set_font(doc.styles["First Paragraph"], font, cfg["body_size_pt"])
    _set_font(doc.styles["Heading 1"], font, cfg["title_size_pt"], True)
    _set_font(doc.styles["Heading 2"], font, cfg["heading_size_pt"], True)
    for name in ("Normal", "Body Text", "First Paragraph"):
        pf = doc.styles[name].paragraph_format
        pf.line_spacing = cfg["line_spacing"]
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
    title = doc.add_paragraph("{{标题}}", style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    sign = doc.add_paragraph(signature, style="First Paragraph")
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sign.paragraph_format.space_after = Pt(6)
    for heading, slot in (("一、现状", "{{现状}}"), ("二、问题和分析", "{{问题和分析}}"), ("三、政策建议", "{{政策建议}}")):
        doc.add_paragraph(heading, style="Heading 2")
        p = doc.add_paragraph(slot, style="First Paragraph")
        p.paragraph_format.first_line_indent = Pt(cfg["body_size_pt"] * 2)
    _apply_run_fonts(doc, font)
    try:
        doc.save(temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def export_submission(template: Path, output: Path, title: str, sections: dict[str, list[str]]) -> None:
    doc = Document(template)
    replacements = {"{{标题}}": title}
    for p in doc.paragraphs:
        if p.text in replacements:
            p.text = replacements[p.text]
            p.style = "Heading 1"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    slot_map = {"{{现状}}": "一、现状", "{{问题和分析}}": "二、问题和分析", "{{政策建议}}": "三、政策建议"}
    for p in list(doc.paragraphs):
        if p.text not in slot_map:
            continue
        key = slot_map[p.text]
        for text in reversed(sections[key]):
            new_p = deepcopy(p._p)
            p._p.addnext(new_p)
            target = next(x for x in doc.paragraphs if x._p is new_p)
            target.text = text
            target.style = "First Paragraph"
        p._element.getparent().remove(p._element)
    _apply_run_fonts(doc, doc.styles["Normal"].font.name or "DengXian")
    output.parent.mkdir(parents=True, exist_ok=True)
    _save_atomic(doc, output)


def parse_submission_markdown(path: Path) -> tuple[str, dict[str, list[str]]]:
    required = ("一、现状", "二、问题和分析", "三、政策建议")
    title = ""
    sections = {name: [] for name in required}
    current: str | None = None
    paragraph: list[str] = []

    def flush() -> None:
        if current and paragraph:
            text = "".join(item.strip() for item in paragraph).strip()
            if text:
                sections[current].append(text)
        paragraph.clear()

    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        heading = line.lstrip("#").strip()
        if heading in required:
            flush()
            current = heading
            continue
        if not line:
            flush()
        elif current:
            paragraph.append(line)
    flush()
    if not title:
        raise ValueError("正式稿 Markdown 缺少一级标题")
    missing = [name for name in required if not sections[name]]
    if missing:
        raise ValueError(f"正式稿 Markdown 缺少内容：{', '.join(missing)}")
    return title, sections
