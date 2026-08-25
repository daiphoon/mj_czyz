from pathlib import Path
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from email.utils import format_datetime
from hashlib import sha256
import json
import shutil
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

from sqmy.config import Settings
from sqmy.document import create_template
from sqmy.evidence import import_evidence_package
from sqmy.workflow import Workflow


class WorkflowTest(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        project_root = Path(__file__).parents[1]
        base = Settings.load(project_root / "config/settings.toml")
        self.settings = Settings(Path(self.tempdir.name), deepcopy(base.raw))
        (self.settings.root / "config").mkdir()
        shutil.copy(project_root / "config/sources.toml", self.settings.root / "config/sources.toml")
        shutil.copy(project_root / "config/policy_mechanisms.toml", self.settings.root / "config/policy_mechanisms.toml")
        cfg = self.settings.section("document")
        create_template(Path(cfg["reference_path"]), self.settings.root / cfg["template_path"], cfg, self.settings.section("project")["signature"])

    def tearDown(self):
        self.tempdir.cleanup()

    def test_mock_flow(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        candidates = wf.scan(run_id)
        self.assertEqual(len(candidates), 5)
        self.assertEqual(wf.scan(run_id)[0].id, candidates[0].id)
        wf.select(run_id, ["C1"])
        wf.pause(run_id, quota=True)
        paused = wf.status(run_id)[0]
        self.assertEqual(paused["status"], "paused_quota")
        self.assertEqual(paused["phase"], "research")
        next_action = wf.resume(run_id)
        self.assertEqual(
            next_action,
            f"sqmy pre-research-check {run_id} CANDIDATE_ID --brief PATH",
        )
        self.assertEqual(wf.status(run_id)[0]["phase"], "research")
        outputs = wf.generate(run_id)
        self.assertEqual(len(outputs), 1)
        self.assertTrue(outputs[0].exists())

    def test_template_uses_dengxian_without_changing_sizes(self):
        template = Document(self.settings.root / self.settings.section("document")["template_path"])
        expected_sizes = {
            "Normal": 12.0,
            "Body Text": 12.0,
            "First Paragraph": 12.0,
            "Heading 1": 20.0,
            "Heading 2": 16.0,
        }
        for style_name, size in expected_sizes.items():
            style = template.styles[style_name]
            self.assertEqual(style.font.name, "DengXian")
            self.assertEqual(style.font.size.pt, size)

    def test_exported_document_sets_dengxian_on_all_visible_runs(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        wf.scan(run_id)
        wf.select(run_id, ["C1"])
        output = wf.generate(run_id)[0]
        document = Document(output)
        visible_runs = [run for paragraph in document.paragraphs for run in paragraph.runs if run.text]
        self.assertTrue(visible_runs)
        for run in visible_runs:
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            self.assertEqual(fonts.get(qn("w:ascii")), "DengXian")
            self.assertEqual(fonts.get(qn("w:hAnsi")), "DengXian")
            self.assertEqual(fonts.get(qn("w:eastAsia")), "DengXian")

    def test_requires_manual_selection(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        with self.assertRaisesRegex(ValueError, "人工确认"):
            wf.generate(run_id)

    def test_candidate_view_does_not_insert_mock_candidates_into_live_run(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run("live")
        self.assertEqual(wf.candidates(run_id), [])
        with wf.db.connect() as conn:
            count = conn.execute(
                "SELECT COUNT(*) FROM candidates WHERE run_id=?", (run_id,)
            ).fetchone()[0]
        self.assertEqual(count, 0)

    def test_selection_rejects_unknown_candidate_without_changing_existing_selection(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        wf.scan(run_id)
        wf.select(run_id, ["C1"])
        with self.assertRaisesRegex(ValueError, "不存在"):
            wf.select(run_id, ["C99"])
        with wf.db.connect() as conn:
            selected = conn.execute(
                "SELECT id FROM candidates WHERE run_id=? AND selected=1", (run_id,)
            ).fetchall()
        self.assertEqual([row["id"] for row in selected], [f"{run_id}:C1"])

    def test_selection_preserves_discovery_checkpoint(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        wf.scan(run_id)
        with wf.db.connect() as conn:
            before = json.loads(
                conn.execute(
                    "SELECT checkpoint_json FROM runs WHERE id=?", (run_id,)
                ).fetchone()[0]
            )
        wf.select(run_id, ["C1"])
        with wf.db.connect() as conn:
            after = json.loads(
                conn.execute(
                    "SELECT checkpoint_json FROM runs WHERE id=?", (run_id,)
                ).fetchone()[0]
            )
        self.assertEqual(after["candidate_file"], before["candidate_file"])
        self.assertEqual(after["selected"], ["C1"])
        self.assertFalse(after["refresh_required"])

    def test_stale_candidate_selection_requires_refresh_but_recent_selection_does_not(self):
        wf = Workflow(self.settings)
        recent = wf.init_run("live")
        wf.scan(recent)
        wf.select(recent, ["C1"])
        recent_status = wf.status(recent)[0]
        self.assertEqual(recent_status["phase"], "research")

        stale = wf.init_run("live")
        wf.scan(stale)
        old = (datetime.now(timezone.utc) - timedelta(hours=25)).isoformat()
        with wf.db.connect() as conn:
            conn.execute("UPDATE runs SET created_at=? WHERE id=?", (old, stale))
        wf.select(stale, ["C1"])
        stale_status = wf.status(stale)[0]
        stale_checkpoint = json.loads(stale_status["checkpoint_json"])
        self.assertEqual(stale_status["phase"], "incremental_review")
        self.assertEqual(stale_checkpoint["next"], f"sqmy refresh {stale}")

    def test_candidate_pool_lists_recent_unselected_live_candidates_only(self):
        wf = Workflow(self.settings)
        live = wf.init_run("live")
        wf.scan(live)
        mock = wf.init_run("mock")
        wf.scan(mock)
        wf.select(live, ["C1"])

        pool = wf.candidate_pool()

        self.assertTrue(pool)
        self.assertTrue(all(item["run_id"] == live for item in pool))
        self.assertNotIn("C1", {item["candidate_id"] for item in pool})

    def test_selection_requires_one_or_two_distinct_candidates(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        wf.scan(run_id)
        with self.assertRaisesRegex(ValueError, "1—2"):
            wf.select(run_id, [])
        with self.assertRaisesRegex(ValueError, "重复"):
            wf.select(run_id, ["C1", "C1"])

    def test_resume_rejects_active_run(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run()
        with self.assertRaisesRegex(ValueError, "只有暂停或失败"):
            wf.resume(run_id)

    def test_refresh_scan_is_zero_model_idempotent_and_requires_human_decision(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run("live")
        wf.scan(run_id)
        wf.select(run_id, ["C1"])
        published = format_datetime(datetime.now(timezone.utc) + timedelta(days=1))
        fixture = self.settings.root / "thursday.json"
        fixture.write_text(json.dumps([{
            "source": {"id": "official", "name": "官方来源", "level": 1, "region": "北京"},
            "xml": (
                "<rss><channel><item><title>北京市教育政策执行问题更新</title>"
                "<link>https://example.gov.cn/thursday</link>"
                "<description>公开数据反映未成年人公共服务问题。</description>"
                f"<pubDate>{published}</pubDate></item></channel></rss>"
            ),
        }], ensure_ascii=False), encoding="utf-8")
        first = wf.incremental_review(run_id, fixture)
        second = wf.incremental_review(run_id, fixture)
        self.assertEqual(first["input_hash"], second["input_hash"])
        self.assertTrue(Path(first["report_path"]).exists())
        with wf.db.connect() as conn:
            calls = conn.execute("SELECT COUNT(*) FROM model_calls WHERE run_id=?", (run_id,)).fetchone()[0]
            tasks = conn.execute(
                "SELECT COUNT(*) FROM tasks WHERE run_id=? AND kind='incremental_review'", (run_id,)
            ).fetchone()[0]
        self.assertEqual(calls, 0)
        self.assertEqual(tasks, 1)
        status = wf.status(run_id)[0]
        self.assertEqual(status["phase"], "research")
        self.assertEqual(status["status"], "pending")
        self.assertIn(
            "incremental_review_decision_next",
            json.loads(status["checkpoint_json"]),
        )
        next_action = wf.record_incremental_decision(run_id, "keep", "未发现替换理由")
        self.assertEqual(
            next_action,
            f"sqmy pre-research-check {run_id} CANDIDATE_ID --brief PATH",
        )
        self.assertEqual(wf.status(run_id)[0]["phase"], "research")
        reused = wf.incremental_review(run_id, fixture)
        self.assertEqual(reused["decision"], "keep")
        resumed_status = wf.status(run_id)[0]
        self.assertEqual(resumed_status["phase"], "research")
        self.assertEqual(resumed_status["status"], "pending")

    def test_thursday_review_does_not_regress_export_state_before_or_after_keep(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run("live")
        wf.scan(run_id)
        wf.select(run_id, ["C1"])
        approve_next = f"sqmy approve topic-{run_id}"
        wf.db.checkpoint(
            run_id,
            phase="export",
            status="needs_review",
            data={"review_path": "/tmp/review.docx", "next": approve_next},
        )
        published = format_datetime(datetime.now(timezone.utc) + timedelta(days=1))
        fixture = self.settings.root / "thursday_export.json"
        fixture.write_text(json.dumps([{
            "source": {"id": "official", "name": "官方来源", "level": 1, "region": "北京"},
            "xml": (
                "<rss><channel><item><title>北京市公共服务政策更新</title>"
                "<link>https://example.gov.cn/thursday-export</link>"
                "<description>公开数据反映公共服务问题。</description>"
                f"<pubDate>{published}</pubDate></item></channel></rss>"
            ),
        }], ensure_ascii=False), encoding="utf-8")

        wf.incremental_review(run_id, fixture)
        pending = wf.status(run_id)[0]
        self.assertEqual(pending["phase"], "export")
        self.assertEqual(pending["status"], "needs_review")
        checkpoint = json.loads(pending["checkpoint_json"])
        self.assertEqual(checkpoint["next"], approve_next)
        self.assertIn("incremental_review_decision_next", checkpoint)

        self.assertEqual(
            wf.record_incremental_decision(run_id, "keep", "未发现足以替换送审稿的新事实或政策"),
            approve_next,
        )
        kept = wf.status(run_id)[0]
        self.assertEqual(kept["phase"], "export")
        self.assertEqual(kept["status"], "needs_review")

    def _import_draftable_evidence(self, topic_id: str) -> None:
        checked_at = datetime.now(timezone.utc).isoformat()
        package = {
            "topic_id": topic_id,
            "sources": [{
                "key": "official",
                "source_name": "政府",
                "page_title": "正式政策",
                "url": f"https://example.gov.cn/{topic_id}",
                "content_hash": "official-v1",
                "source_role": "official_policy",
                "checked_at": checked_at,
            }],
            "claims": [{
                "id": f"{topic_id}-claim",
                "claim_text": "正式政策已经公布相关流程",
                "claim_type": "policy",
                "importance": "critical",
                "epistemic_status": "verified_fact",
                "confidence": "high",
                "uncertainty_reason": "测试仅验证正式政策直接公布的流程",
                "falsifier": "正式政策被撤回、更正或替代",
                "as_of_date": checked_at,
                "sources": [{
                    "source": "official",
                    "role": "supports",
                    "origin_group": "official-policy",
                    "source_level": 1,
                    "primary_source": True,
                }],
            }],
        }
        path = self.settings.root / f"{topic_id}.json"
        path.write_text(json.dumps(package, ensure_ascii=False), encoding="utf-8")
        import_evidence_package(self.settings, path)

    def test_real_draft_requires_evidence_then_approval_before_submission_record(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run("live")
        wf.scan(run_id)
        wf.select(run_id, ["C1"])
        topic_id = "formal-topic"
        source = self.settings.root / "formal.md"
        source.write_text(
            "# 关于完善海淀区测试机制的建议\n\n"
            "中关村支部：李智\n\n"
            "## 一、现状\n\n这是现状。\n\n"
            "## 二、问题和分析\n\n（一）这是问题。\n\n"
            "## 三、政策建议\n\n（一）这是建议。\n",
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "证据闸门"):
            wf.draft(run_id, topic_id, source, candidate_id="C1")
        self._import_draftable_evidence(topic_id)
        with self.assertRaisesRegex(ValueError, "预研闸门"):
            wf.draft(run_id, topic_id, source, candidate_id="C1")
        stamp = datetime.now(timezone.utc).isoformat()
        with wf.db.connect() as conn:
            conn.execute(
                """INSERT INTO research_reviews(
                     id,run_id,candidate_id,topic_id,input_hash,decision,confidence,
                     research_allowed,data_json,report_path,human_decision,human_note,
                     reviewed_at,created_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    "review-1",
                    run_id,
                    "C1",
                    topic_id,
                    "hash",
                    "proceed",
                    "high",
                    1,
                    "{}",
                    "/tmp/review.md",
                    "proceed",
                    "测试确认",
                    stamp,
                    stamp,
                ),
            )
            conn.execute(
                "UPDATE runs SET created_at=? WHERE id=?",
                ((datetime.now(timezone.utc) - timedelta(hours=25)).isoformat(), run_id),
            )
        with self.assertRaisesRegex(ValueError, "来源新鲜度"):
            wf.draft(run_id, topic_id, source, candidate_id="C1")
        refresh_result = {
            "decision": "keep",
            "decision_note": "人工确认无重大变化",
            "decided_at": datetime.now(timezone.utc).isoformat(),
        }
        with wf.db.connect() as conn:
            conn.execute(
                """INSERT INTO tasks(
                     id,run_id,kind,input_hash,status,result_json,token_used,attempts,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?)""",
                (
                    "freshness-review",
                    run_id,
                    "incremental_review",
                    "freshness-hash",
                    "completed",
                    json.dumps(refresh_result, ensure_ascii=False),
                    0,
                    1,
                    datetime.now(timezone.utc).isoformat(),
                ),
            )
        review_path = wf.draft(run_id, topic_id, source, candidate_id="C1")
        self.assertTrue(review_path.exists())
        original_source = source.read_text(encoding="utf-8")
        source.write_text(
            original_source.replace("（一）这是建议。", "（一）这是修订建议。"),
            encoding="utf-8",
        )
        revised_path = wf.draft(run_id, topic_id, source, candidate_id="C1")
        self.assertEqual(revised_path, review_path)
        self.assertIn("这是修订建议", "".join(p.text for p in Document(revised_path).paragraphs))
        with wf.db.connect() as conn:
            export_tasks = conn.execute(
                """SELECT status,result_json FROM tasks
                   WHERE run_id=? AND kind='draft_export' ORDER BY updated_at""",
                (run_id,),
            ).fetchall()
        self.assertEqual([row["status"] for row in export_tasks].count("completed"), 1)
        self.assertEqual([row["status"] for row in export_tasks].count("skipped"), 1)
        completed_result = json.loads(
            next(row["result_json"] for row in export_tasks if row["status"] == "completed")
        )
        self.assertEqual(completed_result["output_sha256"], sha256(revised_path.read_bytes()).hexdigest())

        source.write_text(original_source, encoding="utf-8")
        restored_path = wf.draft(run_id, topic_id, source, candidate_id="C1")
        self.assertNotIn("这是修订建议", "".join(p.text for p in Document(restored_path).paragraphs))
        with wf.db.connect() as conn:
            restored_tasks = conn.execute(
                """SELECT status,attempts FROM tasks
                   WHERE run_id=? AND kind='draft_export'""",
                (run_id,),
            ).fetchall()
        self.assertEqual([row["status"] for row in restored_tasks].count("completed"), 1)
        self.assertEqual([row["status"] for row in restored_tasks].count("skipped"), 1)
        self.assertEqual(max(row["attempts"] for row in restored_tasks), 2)
        with self.assertRaisesRegex(ValueError, "尚未人工通过"):
            wf.mark_submitted(topic_id, "2026-07-20", "海淀区")
        final_path = wf.approve(topic_id)
        self.assertTrue(final_path.exists())
        with wf.db.connect() as conn:
            approved = conn.execute("SELECT * FROM topics WHERE id=?", (topic_id,)).fetchone()
        self.assertEqual(approved["approval_status"], "approved")
        self.assertEqual(approved["actually_submitted"], 0)
        wf.mark_submitted(topic_id, "2026-07-20", "海淀区")
        with wf.db.connect() as conn:
            submitted = conn.execute("SELECT * FROM topics WHERE id=?", (topic_id,)).fetchone()
        self.assertEqual(submitted["actually_submitted"], 1)
        self.assertEqual(submitted["submitted_at"], "2026-07-20")
        self.assertEqual(submitted["submission_level"], "海淀区")
        status = wf.status(run_id)[0]
        checkpoint = json.loads(status["checkpoint_json"])
        self.assertNotIn("warning", checkpoint)
        self.assertEqual(checkpoint["next"], "等待并登记采用情况或反馈")
        self.assertFalse(list(self.settings.root.rglob("*.tmp")))

    def test_run_can_be_explicitly_skipped(self):
        wf = Workflow(self.settings)
        run_id = wf.init_run("live")
        wf.scan(run_id)
        wf.select(run_id, ["C1"])
        wf.skip(run_id, "候选质量不足")
        status = wf.status(run_id)[0]
        self.assertEqual(status["status"], "skipped")
        self.assertIn("候选质量不足", status["checkpoint_json"])
        with wf.db.connect() as conn:
            selected = conn.execute("SELECT COUNT(*) FROM candidates WHERE run_id=? AND selected=1", (run_id,)).fetchone()[0]
        self.assertEqual(selected, 0)
        with self.assertRaisesRegex(ValueError, "运行已关闭"):
            wf.generate(run_id)


if __name__ == "__main__":
    unittest.main()
